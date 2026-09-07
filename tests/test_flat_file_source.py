"""Component A -- flat-file source (.csv / .xlsx).

The file replaces the SQL SELECT and nothing else: sanitization, dedup by
content_hash and the micro-batching all have to behave identically, which is
what most of these assert.
"""

from __future__ import annotations

import csv
from datetime import datetime, timezone

import pytest

pytest.importorskip("pandas")
pytest.importorskip("openpyxl")

from sps.config import IndexingSettings  # noqa: E402
from sps.indexing import (  # noqa: E402
    FlatFileError,
    FlatFileRecordSource,
    IncrementalIndexer,
    Watermark,
)
from sps.indexing.flat_file import build_header_map  # noqa: E402
from sps.vectorstore import InMemoryVectorStore, point_id_for  # noqa: E402
from tests.conftest import TokenOverlapEmbedder  # noqa: E402

# Exactly the columns the evaluation extract carries.
HEADERS = ["SPS_ID", "Part_Number", "Issue_Type", "Problem_Description", "Solution_Text"]

ROWS = [
    ["SPS-1001", "PN-1000", "Quality",
     "Bracket weld seam cracking observed during incoming inspection",
     "Rework the weld seam per the original joint profile and re-inspect."],
    ["SPS-1002", "PN-1000", "Quality",
     "Cracks found in the weld seam of the mounting bracket at goods-in",
     "Segregate the affected lot. Rework the cracked seams and re-inspect."],
    ["SPS-1003", "PN-2000", "Packaging",
     "Outer carton label misprint on shipment packaging",
     "Reprint the carton labels and re-apply before dispatch."],
]

COLD = Watermark.initial()


def write_csv(path, headers=HEADERS, rows=ROWS, encoding="utf-8"):
    with open(path, "w", newline="", encoding=encoding) as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def write_xlsx(path, headers=HEADERS, rows=ROWS, trailing_blank=False):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if trailing_blank:
        sheet.append([None] * len(headers))
        sheet.append([None] * len(headers))
    book.save(path)
    return path


def read_all(path):
    return list(FlatFileRecordSource(path).fetch_since(COLD, 250))


# ------------------------------------------------------------------ parity


def test_csv_and_xlsx_produce_identical_records(tmp_path):
    """Business users supply .xlsx, 300k-row extracts come as .csv. Neither
    format may change what lands in the index."""
    from_csv = read_all(write_csv(tmp_path / "d.csv"))
    from_xlsx = read_all(write_xlsx(tmp_path / "d.xlsx"))

    assert len(from_csv) == len(from_xlsx) == 3
    assert [r.sps_id for r in from_csv] == [r.sps_id for r in from_xlsx]
    assert [r.content_digest() for r in from_csv] == [r.content_digest() for r in from_xlsx]


def test_the_five_specified_columns_map_correctly(tmp_path):
    record = read_all(write_csv(tmp_path / "d.csv"))[0]

    assert record.sps_id == "SPS-1001"
    assert record.part_number == "PN-1000"
    assert record.issue_type == "Quality"
    assert record.problem_description.startswith("Bracket weld seam cracking")
    assert record.actual_solution.startswith("Rework the weld seam")


def test_columns_absent_from_the_extract_become_empty(tmp_path):
    record = read_all(write_csv(tmp_path / "d.csv"))[0]
    assert record.part_description == ""
    assert record.item_status == ""
    assert record.problem_reason_code == ""
    assert record.last_modified_date is None


def test_trailing_blank_rows_are_skipped(tmp_path):
    """Hand-edited workbooks routinely carry empty trailing rows."""
    assert len(read_all(write_xlsx(tmp_path / "d.xlsx", trailing_blank=True))) == 3


# ---------------------------------------------------------------- headers


def test_solution_column_accepts_either_spelling():
    """The SQL query calls it Actual_Solution; the extract calls it
    Solution_Text. Both must work."""
    for header in ("Solution_Text", "Actual_Solution", "Solution"):
        mapping = build_header_map(["SPS_ID", "Problem_Description", header])
        assert "actual_solution" in mapping


def test_headers_match_case_and_separator_insensitively():
    mapping = build_header_map(["sps id", "PROBLEM DESCRIPTION", "solution-text"])
    assert set(mapping) >= {"sps_id", "problem_description", "actual_solution"}


def test_missing_required_column_fails_before_reading_any_row(tmp_path):
    """Better a clear error than 300k records with empty problem descriptions."""
    path = write_csv(
        tmp_path / "d.csv",
        headers=["SPS_ID", "Part_Number", "Issue_Type", "Solution_Text"],
        rows=[["S1", "PN-1", "Quality", "a solution long enough to pass"]],
    )
    with pytest.raises(FlatFileError, match="Problem_Description"):
        read_all(path)


def test_the_error_names_what_it_actually_found(tmp_path):
    path = write_csv(
        tmp_path / "d.csv",
        headers=["Ticket", "Widget"],
        rows=[["S1", "x"]],
    )
    with pytest.raises(FlatFileError) as info:
        read_all(path)
    assert "Ticket" in str(info.value) and "Widget" in str(info.value)


def test_missing_boost_columns_are_warned_not_fatal(caplog):
    with caplog.at_level("WARNING"):
        build_header_map(HEADERS)
    assert "Problem_Reason_Code" in caplog.text
    assert "Last_Modified_Date" in caplog.text


# ------------------------------------------------------------- data fidelity


def test_zero_padded_ids_are_not_turned_into_numbers(tmp_path):
    """pandas would read 00123 as int 123 without dtype=str, silently breaking
    every point ID and every SPS_IDs_Referred citation."""
    path = write_csv(
        tmp_path / "d.csv",
        rows=[["00123", "0045", "Quality", "a" * 30, "b" * 30]],
    )
    record = read_all(path)[0]
    assert record.sps_id == "00123"
    assert record.part_number == "0045"


def test_literal_na_text_is_not_converted_to_nan(tmp_path):
    """A part number of "NA" is data, not a missing value."""
    path = write_csv(
        tmp_path / "d.csv",
        rows=[["S1", "NA", "NULL", "a" * 30, "b" * 30]],
    )
    record = read_all(path)[0]
    assert record.part_number == "NA"
    assert record.issue_type == "NULL"


def test_csv_with_a_utf8_bom_is_read(tmp_path):
    path = write_csv(tmp_path / "d.csv", encoding="utf-8-sig")
    assert len(read_all(path)) == 3


def test_non_ascii_text_survives(tmp_path):
    path = write_csv(
        tmp_path / "d.csv",
        rows=[["S1", "PN-1", "Quality", "Fissure de soudure ≤ 50 µm sur la bride",
               "Reprendre la soudure et contrôler à 10x"]],
    )
    assert "µm" in read_all(path)[0].problem_description


def test_last_modified_date_is_used_when_present(tmp_path):
    path = write_csv(
        tmp_path / "d.csv",
        headers=HEADERS + ["Last_Modified_Date"],
        rows=[["S1", "PN-1", "Quality", "a" * 30, "b" * 30, "2026-03-01T10:00:00+00:00"]],
    )
    record = read_all(path)[0]
    assert record.last_modified_date == datetime(2026, 3, 1, 10, 0, tzinfo=timezone.utc)


def test_an_unparseable_date_is_treated_as_absent(tmp_path):
    path = write_csv(
        tmp_path / "d.csv",
        headers=HEADERS + ["Last_Modified_Date"],
        rows=[["S1", "PN-1", "Quality", "a" * 30, "b" * 30, "last Tuesday"]],
    )
    assert read_all(path)[0].last_modified_date is None


# ------------------------------------------------------------------ errors


def test_missing_file_is_reported_clearly(tmp_path):
    with pytest.raises(FlatFileError, match="not found"):
        FlatFileRecordSource(tmp_path / "nope.csv")


def test_unsupported_extension_is_rejected(tmp_path):
    path = tmp_path / "d.json"
    path.write_text("{}", encoding="utf-8")
    with pytest.raises(FlatFileError, match="Unsupported source file type"):
        FlatFileRecordSource(path)


def test_empty_workbook_is_reported(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    with pytest.raises(FlatFileError):
        read_all(path)


# -------------------------------------------- existing logic is unchanged


def _index(path, tmp_path, store=None, batch_size=250):
    store = store or InMemoryVectorStore()
    report = IncrementalIndexer(
        source=FlatFileRecordSource(path),
        embedder=TokenOverlapEmbedder(),
        store=store,
        settings=IndexingSettings(
            batch_size=batch_size, watermark_path=str(tmp_path / "wm.json")
        ),
    ).run()
    return report, store


def test_short_solutions_are_still_dropped(tmp_path):
    path = write_csv(
        tmp_path / "d.csv",
        rows=ROWS + [["SPS-9", "PN-1", "Quality", "Surface corrosion on the flange face", "TBD"]],
    )
    report, store = _index(path, tmp_path)

    assert report.sanitize.dropped_short_solution == 1
    assert report.indexed == 3
    assert point_id_for("SPS-9") not in store._points


def test_duplicates_are_still_collapsed_by_content_hash(tmp_path):
    duplicate = ["SPS-1099", "PN-1000", "Quality", ROWS[0][3], ROWS[0][4]]
    path = write_csv(tmp_path / "d.csv", rows=ROWS + [duplicate])
    report, store = _index(path, tmp_path)

    assert report.indexed == 3
    assert store.count() == 3
    # Without dates the later SPS_ID wins on id ordering: 1099 > 1001.
    assert point_id_for("SPS-1099") in store._points
    assert point_id_for("SPS-1001") not in store._points


def test_payload_schema_is_unchanged(tmp_path):
    _, store = _index(write_csv(tmp_path / "d.csv"), tmp_path)
    _, payload = store._points[point_id_for("SPS-1001")]
    assert set(payload) == {
        "sps_id", "content_hash", "actual_solution", "part_number",
        "part_description", "item_status", "problem_reason_code", "issue_type",
    }


def test_micro_batching_still_applies(tmp_path):
    rows = [
        [f"S{i:04d}", "PN-1", "Quality", f"Distinct defect {i} on the housing unit",
         f"Fix procedure number {i} applied and verified"]
        for i in range(600)
    ]
    path = write_csv(tmp_path / "big.csv", rows=rows)
    embedder = TokenOverlapEmbedder()
    report = IncrementalIndexer(
        source=FlatFileRecordSource(path),
        embedder=embedder,
        store=InMemoryVectorStore(),
        settings=IndexingSettings(batch_size=250, watermark_path=str(tmp_path / "wm.json")),
    ).run()

    assert report.indexed == 600
    assert report.batches == 3
    assert max(len(call) for call in embedder.passage_calls) <= 250


def test_rerunning_the_same_file_is_idempotent(tmp_path):
    """A flat load is a full load, so a repeat must update in place rather than
    duplicate -- point IDs derive from the SPS_ID, which is what makes that safe."""
    path = write_csv(tmp_path / "d.csv")
    store = InMemoryVectorStore()
    _index(path, tmp_path, store=store)
    assert store.count() == 3

    report, _ = _index(path, tmp_path, store=store)
    assert store.count() == 3
    assert report.indexed == 3  # re-read in full, not skipped by the watermark


def test_watermark_advances_only_when_the_file_carries_dates(tmp_path):
    dated = write_csv(
        tmp_path / "dated.csv",
        headers=HEADERS + ["Last_Modified_Date"],
        rows=[r + ["2026-03-0%d T10:00:00+00:00".replace(" ", "") % (i + 1)]
              for i, r in enumerate(ROWS)],
    )
    report, _ = _index(dated, tmp_path)
    assert report.watermark.last_modified_date.year == 2026

    undated_report, _ = _index(write_csv(tmp_path / "plain.csv"), tmp_path / "other")
    assert undated_report.watermark.last_modified_date.year == 1970


# --------------------------------------------------------------------------
# Numeric-looking identifiers from Excel
# --------------------------------------------------------------------------


def _xlsx_with_part(tmp_path, value):
    """Write a workbook whose Part_Number cell holds `value` at its native type."""
    from openpyxl import Workbook

    path = tmp_path / "numeric.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    sheet.append(["S1", value, "Quality", "a" * 30, "b" * 30])
    book.save(path)
    return path


def test_a_whole_float_part_number_loses_its_decimal(tmp_path):
    """Excel holds every number as a double, so a numeric part number can arrive
    as 1243951.0. Stored verbatim it becomes "1243951.0" -- an identifier that
    matches nothing."""
    record = read_all(_xlsx_with_part(tmp_path, 1243951.0))[0]
    assert record.part_number == "1243951"


def test_an_integer_part_number_is_unchanged(tmp_path):
    assert read_all(_xlsx_with_part(tmp_path, 1243951))[0].part_number == "1243951"


def test_a_hyphenated_part_number_survives(tmp_path):
    """The common house format: text, so leading zeros are safe."""
    assert read_all(_xlsx_with_part(tmp_path, "0012-43951"))[0].part_number == "0012-43951"


def test_leading_zeros_survive_when_the_cell_is_text(tmp_path):
    assert read_all(_xlsx_with_part(tmp_path, "001243951"))[0].part_number == "001243951"


def test_a_genuine_decimal_is_not_truncated(tmp_path):
    """Only whole floats lose the fractional part; real decimals are data."""
    assert read_all(_xlsx_with_part(tmp_path, 12439.51))[0].part_number == "12439.51"


def test_a_float_sps_id_does_not_corrupt_the_point_id(tmp_path):
    """SPS_ID has the same exposure: "1001.0" would be a wrong point ID and a
    wrong citation in SPS_IDs_Referred."""
    from openpyxl import Workbook

    path = tmp_path / "ids.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    sheet.append([1001.0, "PN-1", "Quality", "a" * 30, "b" * 30])
    book.save(path)

    assert read_all(path)[0].sps_id == "1001"


def test_upper_casing_is_a_no_op_for_numeric_part_numbers():
    """Nothing to normalise in the house format, so casing drift cannot occur;
    trimming is the part that does real work."""
    from sps.contracts import normalize_part_number

    for value in ("0012-43951", "001243951", "1243951"):
        assert normalize_part_number(value) == value
    assert normalize_part_number("  0012-43951 ") == "0012-43951"
