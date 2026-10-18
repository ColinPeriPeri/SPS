"""Golden master: the pipeline's observable output, captured and pinned.

Written immediately before the code was consolidated into two files, so it
records what the pipeline did BEFORE that change, and asserts the consolidated
code still does exactly the same.

Only Azure is faked. The intent scorer, the Actor/Judge loop, both
deterministic gates, every prompt builder, the workbook writers and the Tier-2
cache all run for real. The fake client scores deterministically from token
overlap, so the same inputs always reach the same branches.

What is pinned:
  * every cell of the bulk sheet, for history-only and Tier-2 runs
  * every cell of output.xlsx and status.xlsx for a resolved and a refused ticket
  * every message sent to the model, byte for byte
  * the chunks the .docx reader produces from the three sample standards

To re-capture after an INTENDED behaviour change -- never to make a red test
green -- run with SPS_GOLDEN_CAPTURE=1 and review the diff of golden.json.
"""

from __future__ import annotations

import csv
import json
import os
import re
import shutil
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("numpy")
pytest.importorskip("pydantic")

import sps as bulk  # noqa: E402
import sps as resolver  # noqa: E402
from sps import parse_docx  # noqa: E402
from sps import ActorDraft, IntentAssessment, JudgeVerdict  # noqa: E402

GOLDEN = Path(__file__).parent / "golden" / "golden.json"
SAMPLES = Path(__file__).resolve().parent.parent / "samples"
CAPTURE = os.environ.get("SPS_GOLDEN_CAPTURE") == "1"

# The one place the consolidation has to touch: where the chat client is
# looked up. Everything else in this file is layout-independent.
CHAT_CLIENT_TARGET = "sps.AzureOpenAIChatClient"

TIER2_CODE = "RC-EQUIV"
# Columns whose values legitimately differ between runs.
VOLATILE = {"Execution_Timestamp", "Duration_Seconds"}


# ------------------------------------------------------------ the fake model


def _tokens(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", text.lower()))


def _jaccard(a: str, b: str) -> int:
    ta, tb = _tokens(a), _tokens(b)
    return round(100 * len(ta & tb) / len(ta | tb)) if ta | tb else 0


class FakeChatClient:
    """Answers each schema deterministically, and records every call."""

    def __init__(self, log: list) -> None:
        self.log = log

    async def complete_model(self, messages, model):
        self.log.append({"schema": model.__name__, "messages": [dict(m) for m in messages]})
        user = next(m["content"] for m in messages if m["role"] == "user")

        if model is IntentAssessment:
            ticket = user.split("INCOMING TICKET\n", 1)[1].split("\n\nPAST RECORDS", 1)[0]
            blocks = re.findall(r"\[SPS_ID: ([^\]]+)\]\n(.*?)(?=\n\n\[SPS_ID: |\n\Z|\Z)", user, re.S)
            return IntentAssessment(
                ticket_intent="Intent: " + ticket.strip()[:40],
                scores=[
                    {"sps_id": sid, "intent_match": _jaccard(ticket, problem),
                     "reason": f"Token overlap with {sid}."}
                    for sid, problem in blocks
                ],
            )

        if model is ActorDraft:
            label = re.search(r"\[([^\]]+\.docx[^\]]*)\]", user)
            cite = label.group(1) if label else "the extracts"
            # A first draft that the deterministic gate must reject, then a
            # clean rewrite once the critique is in the prompt -- so the
            # gate, the critique wording and the retry prompt are all pinned.
            if "YOUR PREVIOUS DRAFT WAS REJECTED" not in user:
                recommendation = "1. See the feedback in the attachment. 2. Grind out the seam."
            else:
                recommendation = "1. Grind out the cracked seam to sound metal, then re-weld."
            return ActorDraft(
                recommendation=recommendation,
                justification=f"Historical records yielded no resolution. Per {cite}.",
            )

        if model is JudgeVerdict:
            return JudgeVerdict(status="PASS")

        raise AssertionError(f"unexpected schema {model.__name__}")


# --------------------------------------------------------------- helpers


def _read_sheet(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()
    header = [str(h) for h in rows[0]]
    return [
        {h: ("" if v is None else str(v)) for h, v in zip(header, r) if h not in VOLATILE}
        for r in rows[1:]
    ]


def _scrub(value, tmp: Path):
    """Replace run-specific paths so the capture is portable."""
    if isinstance(value, str):
        for form in {str(tmp), tmp.as_posix(), str(tmp).replace("\\", "\\\\")}:
            value = value.replace(form, "<TMP>")
        return value.replace("\\", "/")
    if isinstance(value, list):
        return [_scrub(v, tmp) for v in value]
    if isinstance(value, dict):
        return {k: _scrub(v, tmp) for k, v in value.items()}
    return value


def _write_tickets(path: Path, source: Path, reason_code: str = "") -> Path:
    with open(source, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    columns = list(rows[0]) + (["Problem_Reason_Code"] if reason_code else [])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            if reason_code:
                row["Problem_Reason_Code"] = reason_code
            writer.writerow(row)
    return path


@pytest.fixture
def fake_model(monkeypatch):
    log: list = []
    monkeypatch.setattr(CHAT_CLIENT_TARGET, lambda *a, **k: FakeChatClient(log))
    return log


@pytest.fixture
def no_tier2(monkeypatch):
    monkeypatch.delenv("SPS_TIER2_REASON_CODES", raising=False)


def _capture_or_compare(key: str, actual) -> None:
    store = json.loads(GOLDEN.read_text(encoding="utf-8")) if GOLDEN.exists() else {}
    if CAPTURE:
        store[key] = actual
        GOLDEN.parent.mkdir(parents=True, exist_ok=True)
        GOLDEN.write_text(json.dumps(store, indent=1, ensure_ascii=False, sort_keys=True) + "\n",
                          encoding="utf-8")
        return
    assert key in store, f"no golden capture for {key!r}; run with SPS_GOLDEN_CAPTURE=1"
    assert actual == store[key], f"{key}: output differs from the captured golden master"


# ----------------------------------------------------------------- the runs


@pytest.mark.parametrize(
    "case, tickets, history",
    [
        ("bulk_history_only", "bulk_tickets.csv", "sample_history.csv"),
        ("boilerplate_history_only", "boilerplate_tickets.csv", "boilerplate_history.csv"),
    ],
)
def test_history_only_bulk_run(case, tickets, history, tmp_path, azure_embeddings, fake_model, no_tier2):
    out = tmp_path / "results.xlsx"
    bulk.bulk_main([
        "--tickets", str(SAMPLES / tickets),
        "--history", str(SAMPLES / history),
        "--output", str(out),
    ])
    _capture_or_compare(case, _scrub({"sheet": _read_sheet(out), "calls": fake_model}, tmp_path))


def test_tier2_bulk_run(tmp_path, azure_embeddings, fake_model, monkeypatch):
    """Unknown parts fall through to the standards, gated on reason code."""
    monkeypatch.setenv("SPS_TIER2_REASON_CODES", TIER2_CODE)
    docs = tmp_path / "0250"
    docs.mkdir()
    for doc in sorted((SAMPLES / "0250_docs").glob("*.docx")):
        shutil.copy(doc, docs / doc.name)

    tickets = _write_tickets(tmp_path / "tickets.csv", SAMPLES / "bulk_tickets.csv", TIER2_CODE)
    out = tmp_path / "results.xlsx"
    bulk.bulk_main([
        "--tickets", str(tickets),
        "--history", str(SAMPLES / "sample_history.csv"),
        "--output", str(out),
        "--docs-dir", str(docs),
        # Low enough that the fake embedder's token-overlap scores clear it,
        # so the Tier-2 Actor/Judge loop actually runs rather than gating out.
        "--tier2-threshold", "0.05",
    ])
    _capture_or_compare("bulk_tier2", _scrub({"sheet": _read_sheet(out), "calls": fake_model}, tmp_path))


@pytest.mark.parametrize("case, row_index", [("resolver_first_ticket", 0), ("resolver_unknown_part", 4)])
def test_single_ticket_workbooks(case, row_index, tmp_path, azure_embeddings, fake_model, no_tier2):
    """output.xlsx and status.xlsx exactly as the robot reads them."""
    with open(SAMPLES / "bulk_tickets.csv", newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    ticket = tmp_path / "ticket.csv"
    with open(ticket, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerow(rows[row_index])

    out = tmp_path / "out"
    code = resolver.resolve_main([
        "--ticket-file", str(ticket),
        "--history-file", str(SAMPLES / "sample_history.csv"),
        "--output-dir", str(out),
    ])
    captured = {
        "exit_code": code,
        "status": _read_sheet(out / "status.xlsx"),
        "output": _read_sheet(out / "output.xlsx") if (out / "output.xlsx").exists() else None,
        "calls": fake_model,
    }
    _capture_or_compare(case, _scrub(captured, tmp_path))


def test_docx_reader_chunks():
    """The chunks the .docx reader produces from the three sample standards.

    Captured from python-docx so the stdlib reader that replaces it can be held
    to exactly the same output."""
    chunks = []
    for doc in sorted((SAMPLES / "0250_docs").glob("*.docx")):
        chunks += [
            {"document": c.document, "section": c.section, "text": c.text}
            for c in parse_docx(doc)
        ]
    assert chunks, "no sample .docx parsed"
    _capture_or_compare("docx_chunks", chunks)
