"""Component A -- flat-file source (.csv / .xlsx).

A local-file override for the SQL delta reader, so the index can be built
without database access. Everything downstream is untouched: the same
sanitization, the same content_hash deduplication, the same micro-batching, the
same payload.

Streaming, not loading
----------------------
Both formats are read a chunk at a time. CSV goes through pandas with
`chunksize`; XLSX goes through openpyxl in `read_only` mode, which iterates rows
rather than materialising the sheet. `pandas.read_excel` has no chunked mode and
would pull all 300k rows into memory at once -- exactly the spike that makes CSV
preferable -- so openpyxl is used directly for workbooks. The result is that
either format is safe on the 10-12 GB host.

Full load, not a delta
----------------------
A flat file is a bulk load: every row is read, regardless of the high-water
mark. That is safe to repeat because point IDs are derived from the SPS_ID, so
re-running the same file updates in place rather than duplicating. If the file
carries Last_Modified_Date the watermark still advances, so a later SQL run
picks up only genuinely newer rows.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Iterator, Sequence

from ..contracts import SourceRecord
from .watermark import Watermark

logger = logging.getLogger(__name__)

CSV_SUFFIXES = {".csv", ".txt", ".tsv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}

# Header -> contract field. Several spellings are accepted per field: the SQL
# query calls the solution column Actual_Solution while the extract supplied for
# evaluation calls it Solution_Text, and a file edited by hand may carry either.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "sps_id": ("SPS_ID", "SPSID", "SPS Id", "Id"),
    "problem_description": ("Problem_Description", "Problem Description", "Problem"),
    "actual_solution": (
        "Solution_Text",
        "Actual_Solution",
        "Solution Text",
        "Actual Solution",
        "Solution",
    ),
    "part_number": ("Part_Number", "Part Number", "PartNo"),
    "part_description": ("Part_Description", "Part Description"),
    "item_status": ("Item_Status", "Item Status"),
    "problem_reason_code": ("Problem_Reason_Code", "Problem Reason Code", "Reason_Code"),
    "issue_type": ("Issue_Type", "Issue Type"),
    "last_modified_date": ("Last_Modified_Date", "Last Modified Date", "Modified"),
}

# Without these three a row cannot become an indexable record.
REQUIRED_FIELDS = ("sps_id", "problem_description", "actual_solution")

# Present in the payload schema; absent from a minimal extract. Their loss is
# not fatal but it costs retrieval accuracy, so it is reported once.
BOOSTED_FIELDS = ("part_number", "issue_type", "problem_reason_code")


class FlatFileError(ValueError):
    """The file cannot be used as a record source."""


def _normalise(header: Any) -> str:
    """Fold a header for matching: trim, collapse case and separators."""
    return str(header or "").strip().casefold().replace(" ", "_").replace("-", "_")


def build_header_map(headers: Sequence[Any]) -> dict[str, int]:
    """Map contract field -> column index, matching headers case-insensitively.

    Raises before any row is read if a required column is absent -- far better
    than silently producing 300k records with empty problem descriptions.
    """
    lookup: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = _normalise(header)
        if key and key not in lookup:
            lookup[key] = index

    mapping: dict[str, int] = {}
    for field, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            index = lookup.get(_normalise(alias))
            if index is not None:
                mapping[field] = index
                break

    missing = [f for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        wanted = {f: COLUMN_ALIASES[f][0] for f in missing}
        raise FlatFileError(
            "Source file is missing required column(s): "
            + ", ".join(f"{field} (expected header {header!r})" for field, header in wanted.items())
            + ". Found headers: "
            + ", ".join(repr(str(h)) for h in headers if str(h).strip())
        )

    absent_boosts = [f for f in BOOSTED_FIELDS if f not in mapping]
    if absent_boosts:
        logger.warning(
            "Source file has no %s column(s); those metadata boosts cannot fire, "
            "so composite scores will top out below 1.00 for otherwise perfect matches.",
            ", ".join(COLUMN_ALIASES[f][0] for f in absent_boosts),
        )
    if "last_modified_date" not in mapping:
        logger.warning(
            "Source file has no Last_Modified_Date column. Duplicate problem/solution "
            "pairs will be resolved by SPS_ID order instead of recency, and the "
            "high-water mark will not advance."
        )
    return mapping


def _coerce_date(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    except ValueError:
        logger.debug("Unparseable Last_Modified_Date %r; treating as absent", value)
        return None


def _record_from(row: Sequence[Any], mapping: dict[str, int]) -> SourceRecord:
    def cell(field: str) -> Any:
        index = mapping.get(field)
        if index is None or index >= len(row):
            return ""
        return row[index]

    return SourceRecord.from_row(
        {
            "sps_id": cell("sps_id"),
            "problem_description": cell("problem_description"),
            "actual_solution": cell("actual_solution"),
            "part_number": cell("part_number"),
            "part_description": cell("part_description"),
            "item_status": cell("item_status"),
            "problem_reason_code": cell("problem_reason_code"),
            "issue_type": cell("issue_type"),
            "last_modified_date": _coerce_date(cell("last_modified_date")),
        }
    )


class FlatFileRecordSource:
    """RecordSource backed by a local .csv or .xlsx file."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FlatFileError(f"Source file not found: {self.path}")
        suffix = self.path.suffix.lower()
        if suffix not in CSV_SUFFIXES | EXCEL_SUFFIXES:
            raise FlatFileError(
                f"Unsupported source file type {suffix!r}. "
                f"Expected one of: {', '.join(sorted(CSV_SUFFIXES | EXCEL_SUFFIXES))}"
            )
        self.is_excel = suffix in EXCEL_SUFFIXES

    def fetch_since(self, watermark: Watermark, chunk_size: int) -> Iterator[SourceRecord]:
        """Yield every row in the file.

        The watermark is accepted to satisfy the RecordSource protocol but is
        deliberately not used to filter: a flat file is a bulk load, and
        re-reading it is idempotent because point IDs derive from the SPS_ID.
        """
        logger.info(
            "Reading %s source %s (full load; watermark not applied)",
            "Excel" if self.is_excel else "CSV",
            self.path,
        )
        reader = self._read_excel if self.is_excel else self._read_csv
        yielded = 0
        for record in reader(chunk_size):
            yielded += 1
            yield record
        logger.info("Read %d row(s) from %s", yielded, self.path)

    # -- CSV ---------------------------------------------------------------

    def _read_csv(self, chunk_size: int) -> Iterator[SourceRecord]:
        import pandas as pd

        separator = "\t" if self.path.suffix.lower() == ".tsv" else ","
        mapping: dict[str, int] | None = None
        # dtype=str keeps a zero-padded SPS_ID like "00123" from becoming 123;
        # keep_default_na=False keeps a literal "NA" or "NULL" as text rather
        # than NaN, so sanitization sees what the file actually says.
        frames = pd.read_csv(
            self.path,
            chunksize=max(chunk_size, 1),
            dtype=str,
            keep_default_na=False,
            sep=separator,
            encoding="utf-8-sig",  # tolerate a BOM, as with every other input
        )
        for frame in frames:
            if mapping is None:
                mapping = build_header_map(list(frame.columns))
            for row in frame.itertuples(index=False, name=None):
                yield _record_from(row, mapping)
            del frame

    # -- Excel -------------------------------------------------------------

    def _read_excel(self, chunk_size: int) -> Iterator[SourceRecord]:
        from openpyxl import load_workbook

        # read_only streams rows instead of building the whole sheet in memory;
        # data_only returns computed values rather than formula strings, which
        # matters when a business user has built the extract with formulas.
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = next(rows)
            except StopIteration:
                raise FlatFileError(f"Source file is empty: {self.path}") from None
            mapping = build_header_map(list(headers))
            for row in rows:
                if row is None or all(cell in (None, "") for cell in row):
                    continue  # trailing blank rows are common in hand-edited sheets
                yield _record_from(row, mapping)
        finally:
            workbook.close()
