"""Regenerate the demonstration 0250 documents under samples/0250_docs/.

The .docx files are committed, so this is only needed when the demo content
changes. They exist so Tier 2 can be exercised end to end before any real
standard is loaded -- and so `samples/eval_cases/case04`, a defect with no
historical precedent, has somewhere to land.

    python -m scripts.make_sample_docs [--output-dir samples/0250_docs]

The content is invented. It is shaped like a real standard -- numbered
sections, a limits table, a disposition that names its own re-inspection step
-- because the point is to exercise the parser's section handling and the
Actor's grounding check, and neither is exercised by lorem ipsum.
"""

from __future__ import annotations

import argparse
from pathlib import Path

DEFAULT_OUTPUT = Path("samples/0250_docs")

WELD = [
    ("4.1 Scope", [
        "This standard defines acceptance criteria and supplier disposition for "
        "welded assemblies supplied against a drawing that invokes it. It applies "
        "to fillet and butt welds on carbon steel and stainless steel brackets, "
        "mounts and frames."
    ]),
    ("4.2 Weld Seam Cracking", [
        "Cracking in a fillet or butt weld seam is cause for rejection of the "
        "affected part. Transverse cracks, longitudinal cracks and crater cracks "
        "shall all be treated as rejectable regardless of length. On discovery at "
        "incoming inspection, segregate the affected lot and quarantine it "
        "pending disposition.",
        "The cracked seam shall be ground out to sound metal for the full length "
        "of the crack plus 25 mm beyond each visible end. Re-weld to the original "
        "joint profile using the qualified procedure for the base material. After "
        "re-welding, the repaired region shall be re-inspected by dye penetrant "
        "in accordance with section 7.1.",
        "A part that fails re-inspection twice shall be scrapped and shall not be "
        "reworked a third time. Record each repair against the lot traveller and "
        "return the traveller with the shipment.",
    ]),
    ("4.3 Weld Porosity Limits", [
        "Porosity shall be assessed visually at 10x magnification over the full "
        "length of the seam. The limits below apply to every weld covered by this "
        "standard.",
        "__TABLE__",
        "Where porosity exceeds the limit, grind out the porous region to sound "
        "metal and re-weld to the original joint profile. Verify shielding gas "
        "flow rate and base material cleanliness before the next production run, "
        "and state the corrective action taken on the certificate of conformance.",
    ]),
    ("4.4 Weld Undercut", [
        "Undercut deeper than 0.5 mm, or exceeding 10 percent of the base "
        "material thickness where that is less, is rejectable. Undercut within "
        "the limit requires no rework and shall not be blended.",
    ]),
    ("7.1 Re-inspection After Repair", [
        "Every repaired weld shall be re-inspected before the part is released. "
        "Dye penetrant inspection is the default method. Radiographic inspection "
        "is required where the joint is a full penetration butt weld carrying a "
        "primary load path. Re-inspection results shall accompany the shipment.",
    ]),
]

POROSITY_TABLE = [
    ("Condition", "Limit"),
    ("Scattered porosity", "2 percent by area maximum over any 25 mm of seam"),
    ("Clustered porosity", "Rejectable at any level"),
    ("Wormholes and blowholes", "Rejectable at any level"),
    ("Crater pipe", "Rejectable where it breaks the surface"),
]

PACKAGING = [
    ("2.1 Carton Labelling", [
        "Each outer carton shall carry a label showing the part number, the "
        "drawing revision, the quantity contained, the purchase order number and "
        "the date of packing. Labels shall be printed at 300 dpi or better and "
        "shall remain legible after transit.",
        "Where a label is found unreadable on arrival at the receiving dock, the "
        "carton shall be held and the supplier notified. Reprint the labels, "
        "apply them over the originals without obscuring any other marking, and "
        "verify print quality against a retained sample before release.",
    ]),
    ("2.4 Barcode Symbology", [
        "Barcodes shall use Code 128 unless the purchase order specifies "
        "otherwise. The encoded value shall match the printed part number "
        "exactly, including leading zeros and the dash. A barcode that scans to a "
        "different value than the printed text is a rejectable nonconformance and "
        "the whole shipment shall be held.",
    ]),
]

# A standard that states a requirement but no disposition. Its job in the demo
# set is to be retrieved and then declined: an Actor that answers from this has
# invented the corrective action, which is the failure mode Tier 2's grounding
# check exists to catch.
SURFACE = [
    ("3.2 Surface Finish", [
        "Machined surfaces called out on the drawing shall achieve a surface "
        "roughness of Ra 1.6 micrometres or better. Surfaces not called out shall "
        "be free of burrs, raised edges and loose swarf. Roughness shall be "
        "measured with a calibrated profilometer over a 12 mm evaluation length, "
        "taking the arithmetic mean of three readings spaced along the surface.",
    ]),
]


def build(path: Path, title: str, sections, table=None) -> Path:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=2)
        for text in paragraphs:
            if text == "__TABLE__" and table:
                grid = document.add_table(rows=len(table), cols=2)
                grid.style = "Table Grid"
                for r, (left, right) in enumerate(table):
                    grid.cell(r, 0).text = left
                    grid.cell(r, 1).text = right
                continue
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m scripts.make_sample_docs")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args(argv)
    out = Path(args.output_dir)

    written = [
        build(out / "0250-Weld-Standards.docx", "0250 Weld Standards",
              WELD, POROSITY_TABLE),
        build(out / "0250-Packaging-Standards.docx", "0250 Packaging Standards",
              PACKAGING),
        build(out / "0250-Surface-Finish.docx", "0250 Surface Finish",
              SURFACE),
    ]
    for path in written:
        print(f"{path}  {path.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
